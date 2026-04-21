from docx import Document
from openpyxl import Workbook
def save_data(data, format):
    if format == "doc":
        doc = Document(); doc.add_paragraph(str(data)); doc.save("out.docx")
    else:
        wb = Workbook(); wb.active.append([str(data)]); wb.save("out.xlsx")